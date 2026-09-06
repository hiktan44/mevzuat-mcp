"""Ticaret politikası savunma araçları, gözetim ve İthalat Tebliğleri: resmî listelerden GTİP eşlemesi
ve günlük otomatik eşitleme.

Kaynaklar:
- Dampinge/sübvansiyona karşı kesin önlemler: Ticaret Bakanlığı İthalat Genel Müdürlüğü
  "Yürürlükteki Önlemler" çalışma kitabı (damping sayfasındaki .xlsx bağlantısı her güncellemede
  yeni tarihle yayımlanır; bağlantı sayfadan keşfedilir).
- Korunma önlemleri: Korunma Önlemleri Dairesi "Yürürlükte Bulunan Korunma Önlemleri" çalışma kitabı.
- Gözetim: mevzuat.gov.tr'de yürürlükteki "İthalatta Gözetim Uygulanmasına İlişkin Tebliğ" metinleri
  (resmî arama + fihrist metni; GTİP / birim gümrük kıymeti tabloları).
- İthalat Tebliğleri dizini: Ticaret Bakanlığı "İthalat Tebliğleri (yıl)" sayfası.

Eşleme kuralı: resmî satırdaki kod, sorgulanan GTİP'in ön eki ise (4, 6, 8, 10 veya 12 hane) satır
kapsam adayıdır. Damping önlemlerinde menşe ülke de eşleşmelidir; korunma önlemleri menşe ayrımı
olmadan (ülkeye özgü olanlar hariç) uygulanır; gözetim menşeden bağımsızdır. Önlem oranı resmî
tabloda yazıldığı gibi metin olarak döner; hesaba otomatik girilmez, kullanıcı doğrulaması istenir.

Depo: `data/official/*.json` tohum dosyaları ilk açılışta yüklenir; günlük eşitleme sonuçları
`trade_measures.sqlite3` içinde saklanır ve her fark değişiklik defterine yazılır.
"""

from __future__ import annotations

import asyncio
import io
import json
import logging
import os
import re
import sqlite3
import time
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import quote, urljoin

import httpx
from bs4 import BeautifulSoup

from countries import find_country
from trusted_certificates import GEOTRUST_TLS_RSA_CA_G1_PEM

logger = logging.getLogger(__name__)

ROOT = Path(__file__).resolve().parent
SEED_DIR = ROOT / "data" / "official"
DEFAULT_SYNC_INTERVAL = int(os.environ.get("TRADE_MEASURES_SYNC_INTERVAL_SECONDS", "86400"))
USER_AGENT = "Mozilla/5.0 (compatible; MevzuatMCP/1.5; +https://gumruksor.com/)"

ANTIDUMPING_PAGE = "https://ticaret.gov.tr/ithalat/ticaret-politikasi-savunma-araclari/damping-ve-subvansiyon"
SAFEGUARD_PAGE = "https://ticaret.gov.tr/ithalat/ticaret-politikasi-savunma-araclari/korunma-onlemleri/yururlukteki-onlemler"
COMMUNIQUE_PAGE = (
    "https://ticaret.gov.tr/ithalat/ithalat-mevzuati/ithalat-rejimi-karari-igv-karari-ve-ithalat-tebligleri/"
    "3-ithalat-tebligleri-{year}-yili"
)
MEVZUAT_HOME = "https://www.mevzuat.gov.tr/"
MEVZUAT_DATATABLE = "https://www.mevzuat.gov.tr/Anasayfa/MevzuatDatatable"
MEVZUAT_IFRAME = "https://www.mevzuat.gov.tr/anasayfa/MevzuatFihristDetayIframe?MevzuatTur=9&MevzuatNo={no}&MevzuatTertip=5"
MEVZUAT_PUBLIC = "https://www.mevzuat.gov.tr/mevzuat?MevzuatNo={no}&MevzuatTur=9&MevzuatTertip=5"
SURVEILLANCE_TITLE = "İthalatta Gözetim Uygulanmasına İlişkin Tebliğ"

_CODE_RE = re.compile(r"\d{4}(?:\.\d{2}){0,4}|\d{2}\.\d{2}(?:\.\d{2}){0,4}|\d{4,12}")
_ROW_CODE_RE = re.compile(r"^\s*(\d{4}(?:\.\d{2}){1,4}|\d{6,12})\s*$")
_ALL_COUNTRIES = {"tüm ülkeler", "tum ulkeler", "all countries"}
KINDS = ("anti_dumping", "safeguard", "surveillance", "communiques")
KIND_LABELS = {
    "anti_dumping": "Damping / sübvansiyon önlemleri",
    "safeguard": "Korunma önlemleri",
    "surveillance": "Gözetim tebliğleri",
    "communiques": "İthalat Tebliğleri dizini",
}


# --------------------------------------------------------------------------- helpers

def normalise_code(value: str) -> str:
    """'7306.40.20.90.00' -> '730640209000'; '55.13' -> '5513'."""
    return re.sub(r"\D", "", value or "")


def extract_codes(raw: str) -> list[str]:
    """Resmî hücrelerdeki serbest metinden GTİP/GTP kodlarını (rakam dizisi) çıkarır."""
    codes: list[str] = []
    for token in _CODE_RE.findall(raw or ""):
        digits = normalise_code(token)
        if len(digits) < 4 or len(digits) > 12 or len(digits) % 2:
            continue
        if digits not in codes:
            codes.append(digits)
    return codes


def code_matches(candidate: str, gtip: str) -> bool:
    """Resmî satır kodu sorgulanan kodun ön eki ise (ya da tersi, kısa sorgularda) kapsamdadır."""
    if not candidate or not gtip:
        return False
    return gtip.startswith(candidate) or candidate.startswith(gtip)


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return re.sub(r"[ \t\xa0]+", " ", str(value)).strip()


def _country_key(value: str) -> str | None:
    if not value:
        return None
    country = find_country(value.split("/")[0].strip())
    return country.key if country else None


def _country_matches(origin: str | None, measure_country: str) -> bool | None:
    """True/False eşleşme; menşe verilmediyse None (bilinmiyor)."""
    if not origin:
        return None
    lowered = measure_country.strip().lower()
    if lowered in _ALL_COUNTRIES or lowered.startswith("tüm ülkeler") or lowered.startswith("tum ulkeler"):
        return True
    origin_key = _country_key(origin)
    measure_key = _country_key(measure_country)
    if origin_key and measure_key:
        return origin_key == measure_key
    return origin.strip().lower() == measure_country.split("/")[0].strip().lower()


def _parse_unit_value(text: str) -> float | None:
    cleaned = (text or "").strip()
    if not cleaned or cleaned in {"-", "—"}:
        return None
    match = re.search(r"\d+(?:[.,]\d+)?", cleaned)
    if not match:
        return None
    number = match.group().replace(",", ".")
    try:
        return float(number)
    except ValueError:
        return None


def _status(expires: str | None, today: date) -> str:
    if not expires:
        return "unknown"
    try:
        return "in_force" if date.fromisoformat(expires[:10]) >= today else "expired"
    except ValueError:
        return "unknown"


def official_ssl_context() -> Any:
    """Resmî sunucular ara sertifikayı göndermediği için yalnız o ara CA eklenir."""
    context = httpx.create_ssl_context()
    context.load_verify_locations(cadata=GEOTRUST_TLS_RSA_CA_G1_PEM)
    return context


# --------------------------------------------------------------------------- parsers (pure)

def parse_antidumping_workbook(payload: bytes) -> dict[str, list[dict[str, str]]]:
    """Bakanlığın 'Yürürlükteki Önlemler' kitabından kesin ve geçici önlem satırları."""
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    out: dict[str, list[dict[str, str]]] = {"definitive": [], "provisional": []}
    for sheet in workbook.worksheets:
        title = sheet.title.lower()
        if "kesin" in title or "definitive" in title:
            key = "definitive"
        elif "geçici" in title or "gecici" in title or "prov" in title:
            key = "provisional"
        else:
            continue
        header_seen = False
        for row in sheet.iter_rows(values_only=True):
            values = [_clean(cell) for cell in row]
            if not header_seen:
                header_seen = bool(values) and values[0] == "DOSYA NO"
                continue
            if not values or not values[0] or values[0] == "CASE NUMBER":
                continue
            values += [""] * (40 - len(values))
            if key == "definitive":
                out[key].append(
                    {
                        "file_no": values[0], "sector": values[1], "product": values[2], "product_en": values[3],
                        "gtip": values[4], "country": values[5], "country_en": values[6], "communique": values[7],
                        "rg_date": values[8], "rg_no": values[9], "rate": values[11], "kind": values[12],
                        "expires": values[13], "orig_initiation": values[14], "orig_provisional": values[15],
                        "orig_definitive": values[16], "expiry_review_open": values[20], "expiry_review_close": values[21],
                        "expiry_review_result": values[22], "circumvention_open": values[26],
                        "circumvention_definitive": values[28], "notes": " | ".join(v for v in values[33:37] if v),
                    }
                )
            else:
                out[key].append(
                    {
                        "file_no": values[0], "sector": values[1], "product": values[2], "product_en": values[3],
                        "gtip": values[4], "country": values[5], "country_en": values[6], "communique": values[7],
                        "rg_date": values[8], "rg_no": values[9], "rate": values[10], "kind": values[11],
                        "initiation": values[12],
                    }
                )
    if not out["definitive"]:
        raise ValueError("Damping çalışma kitabında kesin önlem satırı bulunamadı; biçim değişmiş olabilir.")
    return out


def parse_safeguard_workbook(payload: bytes) -> list[dict[str, Any]]:
    """Korunma Önlemleri Dairesi listesi: her önlem, alt satırlarında işlem adımları ve tutarları."""
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    measures: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for row in sheet.iter_rows(values_only=True):
        values = [_clean(cell) for cell in row]
        values += [""] * (18 - len(values))
        if values[1].isdigit() and values[2]:
            current = {
                "seq": int(values[1]), "file_no": values[2], "product": values[3], "gtip": values[4],
                "country": values[5], "stage": values[12], "original_start": values[13], "expires": values[14],
                "amounts": [], "acts": [],
            }
            measures.append(current)
        if current is None:
            continue
        if values[6]:
            current["acts"].append(
                {"kind": values[6], "number": values[7], "rg_no": values[8], "rg_date": values[9], "published": values[10], "wto": values[11]}
            )
        if values[15] and values[15] != "-":
            current["amounts"].append(values[15])
    if not measures:
        raise ValueError("Korunma önlemi listesinde satır bulunamadı; biçim değişmiş olabilir.")
    return measures


def _header_unit(text: str) -> str | None:
    match = re.search(r"\(([^)]*)\)", text)
    if match:
        return match.group(1).replace("*", "").strip()
    match = re.search(r"(ABD Doları\s*/\s*\w+|USD\s*/\s*\w+|\$\s*/\s*\w+)", text, flags=re.I)
    return match.group(1).replace("*", "").strip() if match else None


def parse_surveillance_page(html_text: str, meta: dict[str, Any] | None = None) -> dict[str, Any]:
    """Gözetim tebliği metnindeki GTİP / eşya tanımı / birim gümrük kıymeti tablolarını çıkarır."""
    soup = BeautifulSoup(html_text, "html.parser")
    items: list[dict[str, Any]] = []
    default_unit: str | None = None
    for table in soup.find_all("table"):
        rows = table.find_all("tr")
        if not rows:
            continue
        gtip_col = 0
        desc_col: int | None = None
        value_col: int | None = None
        country_col: int | None = None
        unit: str | None = None
        header_found = False
        for tr in rows:
            cells = [re.sub(r"\s+", " ", td.get_text(" ", strip=True)) for td in tr.find_all(["td", "th"])]
            if not cells:
                continue
            joined = " ".join(cells).lower()
            if not header_found and ("g.t.i.p" in joined or "gtip" in joined or "gtİp" in joined or "g.t.İ.p" in joined or "gtp" in joined):
                header_found = True
                for index, cell in enumerate(cells):
                    low = cell.lower()
                    if any(token in low for token in ("g.t.i.p", "g.t.İ.p", "gtip", "gtİp", "gtp")) and "tanım" not in low:
                        gtip_col = index
                    elif "tanım" in low or "eşya" in low or "esya" in low or "madde" in low:
                        desc_col = index
                    elif "kıymet" in low or "kiymet" in low or "abd dolar" in low or "$" in low or "usd" in low or "fiyat" in low:
                        value_col = index
                        unit = _header_unit(cell) or unit
                    elif "menşe" in low or "mense" in low or "ülke" in low:
                        country_col = index
                if unit:
                    default_unit = unit
                continue
            code_cell = cells[gtip_col] if gtip_col < len(cells) else ""
            if not _ROW_CODE_RE.match(code_cell):
                # Bazı tablolarda kod ile açıklama aynı hücrede: "8481.10.05.00.00 Filtre..."
                match = re.match(r"^\s*(\d{4}(?:\.\d{2}){1,4})\s+(.+)$", code_cell)
                if not match:
                    continue
                code_text, inline_desc = match.group(1), match.group(2)
            else:
                code_text, inline_desc = code_cell.strip(), ""
            description = inline_desc
            if desc_col is not None and desc_col < len(cells):
                description = cells[desc_col] or description
            elif len(cells) > gtip_col + 1 and value_col is None:
                description = cells[gtip_col + 1]
            value = ""
            if value_col is not None and value_col < len(cells):
                value = cells[value_col]
            else:
                numeric = [c for c in cells[gtip_col + 1 :] if re.fullmatch(r"[\d.,]+", c.replace(" ", ""))]
                value = numeric[-1] if numeric else ""
            item: dict[str, Any] = {
                "gtip": code_text,
                "description": description,
                "value": value,
                "unit": unit or default_unit,
            }
            if country_col is not None and country_col < len(cells):
                item["country"] = cells[country_col]
            items.append(item)
    text = soup.get_text(" ", strip=True)
    if default_unit is None:
        found = re.search(r"\(ABD Doları\s*/\s*([A-Za-zÇĞİÖŞÜçğıöşü0-9]+)\*?\)", text)
        if found:
            default_unit = f"ABD Doları/{found.group(1)}"
            for item in items:
                item["unit"] = item["unit"] or default_unit
    title_match = re.search(r"(İTHALATTA GÖZETİM UYGULANMASINA İLİŞKİN TEBLİĞ\s*\(TEBLİĞ NO:\s*[^)]+\))", text, flags=re.I)
    result = dict(meta or {})
    result.update(
        {
            "title": result.get("title") or (title_match.group(1).strip() if title_match else SURVEILLANCE_TITLE),
            "unit": default_unit,
            "items": items,
            "item_count": len(items),
        }
    )
    return result


def discover_workbook_link(html_text: str, base_url: str, *, must_contain: Iterable[str]) -> str | None:
    """Sayfadaki .xls/.xlsx bağlantıları arasından metni/adresi ipuçlarını içeren ilkini döndürür."""
    soup = BeautifulSoup(html_text, "html.parser")
    hints = [hint.lower() for hint in must_contain]
    for anchor in soup.find_all("a", href=True):
        href = anchor["href"]
        text = anchor.get_text(" ", strip=True).lower()
        if not re.search(r"\.xlsx?(\?|$)", href, flags=re.I):
            continue
        haystack = f"{text} {href.lower()}"
        if all(hint in haystack for hint in hints):
            return urljoin(base_url, href.replace(" ", "%20"))
    return None


def parse_communique_index(html_text: str, base_url: str, year: int) -> list[dict[str, str]]:
    """Bakanlığın 'İthalat Tebliğleri (yıl)' sayfasındaki tebliğ bağlantıları."""
    soup = BeautifulSoup(html_text, "html.parser")
    entries: list[dict[str, str]] = []
    seen: set[str] = set()
    for anchor in soup.find_all("a", href=True):
        title = re.sub(r"\s+", " ", anchor.get_text(" ", strip=True)).strip("–—- ")
        href = urljoin(base_url, anchor["href"])
        if not title or href in seen:
            continue
        if "tebliğ" not in title.lower() and "teblig" not in title.lower():
            continue
        if not re.search(r"resmigazete|mevzuat\.gov|ticaret\.gov\.tr/data|\.pdf", href, flags=re.I):
            continue
        number = re.search(r"\(İthalat:\s*(\d{4}/\d+)\)", title, flags=re.I)
        seen.add(href)
        entries.append(
            {
                "title": title,
                "url": href,
                "number": number.group(1) if number else "",
                "year": str(year),
                "gazette": re.search(r"/eskiler/(\d{4})/(\d{2})/(\d{8})", href).group(3) if re.search(r"/eskiler/(\d{4})/(\d{2})/(\d{8})", href) else "",
            }
        )
    return entries


# --------------------------------------------------------------------------- store

@dataclass(frozen=True)
class MeasureHit:
    measure_type: str  # anti_dumping | countervailing | safeguard | surveillance
    matched_code: str
    country: str
    origin_match: bool | None
    rate_text: str
    unit_value_usd: float | None
    unit: str | None
    product: str
    legal_act: str
    gazette: str
    expires: str | None
    status: str  # in_force | expired | unknown
    notes: str = ""
    source: str = ""

    def as_dict(self) -> dict[str, Any]:
        return {
            "measure_type": self.measure_type, "matched_code": self.matched_code, "country": self.country,
            "origin_match": self.origin_match, "rate_text": self.rate_text, "unit_value_usd": self.unit_value_usd,
            "unit": self.unit, "product": self.product, "legal_act": self.legal_act, "gazette": self.gazette,
            "expires": self.expires, "status": self.status, "notes": self.notes, "source": self.source,
        }


@dataclass
class TradeMeasureReport:
    gtip: str
    origin_country: str | None
    as_of: str
    anti_dumping: list[MeasureHit] = field(default_factory=list)
    safeguard: list[MeasureHit] = field(default_factory=list)
    surveillance: list[MeasureHit] = field(default_factory=list)
    sources: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def surveillance_unit_value(self) -> float | None:
        values = [hit.unit_value_usd for hit in self.surveillance if hit.unit_value_usd and hit.status == "in_force"]
        return max(values) if values else None

    @property
    def applicable_anti_dumping(self) -> list[MeasureHit]:
        return [hit for hit in self.anti_dumping if hit.origin_match is not False and hit.status != "expired"]

    def as_dict(self) -> dict[str, Any]:
        return {
            "gtip": self.gtip, "origin_country": self.origin_country, "as_of": self.as_of,
            "anti_dumping": [hit.as_dict() for hit in self.anti_dumping],
            "safeguard": [hit.as_dict() for hit in self.safeguard],
            "surveillance": [hit.as_dict() for hit in self.surveillance],
            "surveillance_unit_value": self.surveillance_unit_value,
            "sources": self.sources, "warnings": self.warnings,
        }


class TradeMeasureStore:
    """Tohum JSON + sqlite anlık görüntüleri; her kayıt farkı değişiklik defterine yazılır."""

    def __init__(self, data_dir: str | Path | None = None, seed_dir: str | Path | None = None) -> None:
        root = Path(data_dir or os.environ.get("MEVZUAT_DATA_DIR") or ROOT)
        root.mkdir(parents=True, exist_ok=True)
        self.db_path = root / "trade_measures.sqlite3"
        self.seed_dir = Path(seed_dir or SEED_DIR)
        self._cache: dict[str, Any] = {}
        self._init_schema()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.db_path, timeout=30)
        connection.row_factory = sqlite3.Row
        return connection

    def _init_schema(self) -> None:
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS datasets (
                    kind TEXT PRIMARY KEY, fetched_at TEXT NOT NULL, source_url TEXT NOT NULL,
                    source_label TEXT NOT NULL, item_count INTEGER NOT NULL, payload TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS surveillance_docs (
                    mevzuat_no TEXT PRIMARY KEY, title TEXT NOT NULL, rg_date TEXT, rg_no TEXT, url TEXT,
                    fetched_at TEXT NOT NULL, payload TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS changes (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, kind TEXT NOT NULL, changed_at TEXT NOT NULL,
                    added INTEGER NOT NULL, removed INTEGER NOT NULL, modified INTEGER NOT NULL, detail TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_changes_kind ON changes(kind, changed_at);
                """
            )

    # ---- seeds
    def _seed(self, kind: str) -> Any:
        names = {
            "anti_dumping": "antidumping_measures.json",
            "safeguard": "safeguard_measures.json",
            "surveillance": "surveillance_measures.json",
            "communiques": "import_communiques.json",
        }
        path = self.seed_dir / names[kind]
        if not path.exists():
            return None
        with path.open(encoding="utf-8") as handle:
            return json.load(handle)

    def load(self, kind: str) -> Any:
        if kind in self._cache:
            return self._cache[kind]
        with self._connect() as connection:
            row = connection.execute("SELECT payload FROM datasets WHERE kind=?", (kind,)).fetchone()
        payload = json.loads(row["payload"]) if row else self._seed(kind)
        self._cache[kind] = payload
        return payload

    def metadata(self, kind: str) -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT fetched_at, source_url, source_label, item_count FROM datasets WHERE kind=?", (kind,)
            ).fetchone()
        if row:
            return {"origin": "synced", **dict(row)}
        seed = self._seed(kind)
        if seed is None:
            return {"origin": "missing", "fetched_at": None, "source_url": None, "source_label": None, "item_count": 0}
        return {
            "origin": "seed",
            "fetched_at": None,
            "source_url": None,
            "source_label": "depo tohum dosyası (data/official)",
            "item_count": _count_items(kind, seed),
        }

    def save(self, kind: str, payload: Any, *, source_url: str, source_label: str) -> dict[str, Any]:
        previous = self.load(kind)
        diff = diff_payloads(kind, previous, payload)
        now = datetime.now(UTC).isoformat(timespec="seconds")
        with self._connect() as connection:
            connection.execute(
                "INSERT INTO datasets(kind,fetched_at,source_url,source_label,item_count,payload) VALUES(?,?,?,?,?,?) "
                "ON CONFLICT(kind) DO UPDATE SET fetched_at=excluded.fetched_at, source_url=excluded.source_url, "
                "source_label=excluded.source_label, item_count=excluded.item_count, payload=excluded.payload",
                (kind, now, source_url, source_label, _count_items(kind, payload), json.dumps(payload, ensure_ascii=False)),
            )
            if previous is not None and diff["total"]:
                connection.execute(
                    "INSERT INTO changes(kind,changed_at,added,removed,modified,detail) VALUES(?,?,?,?,?,?)",
                    (kind, now, diff["added_count"], diff["removed_count"], diff["modified_count"], json.dumps(diff, ensure_ascii=False)),
                )
        self._cache[kind] = payload
        return diff

    # ---- surveillance document cache
    def surveillance_doc(self, mevzuat_no: str) -> dict[str, Any] | None:
        with self._connect() as connection:
            row = connection.execute("SELECT payload FROM surveillance_docs WHERE mevzuat_no=?", (mevzuat_no,)).fetchone()
        return json.loads(row["payload"]) if row else None

    def save_surveillance_doc(self, doc: dict[str, Any]) -> None:
        with self._connect() as connection:
            connection.execute(
                "INSERT OR REPLACE INTO surveillance_docs(mevzuat_no,title,rg_date,rg_no,url,fetched_at,payload) VALUES(?,?,?,?,?,?,?)",
                (
                    str(doc["mevzuat_no"]), doc.get("title", ""), doc.get("rg_date", ""), doc.get("rg_no", ""), doc.get("url", ""),
                    datetime.now(UTC).isoformat(timespec="seconds"), json.dumps(doc, ensure_ascii=False),
                ),
            )

    def changes(self, kind: str | None = None, *, limit: int = 20) -> list[dict[str, Any]]:
        query = "SELECT id, kind, changed_at, added, removed, modified, detail FROM changes"
        params: tuple[Any, ...] = ()
        if kind:
            query += " WHERE kind=?"
            params = (kind,)
        query += " ORDER BY id DESC LIMIT ?"
        with self._connect() as connection:
            rows = connection.execute(query, (*params, limit)).fetchall()
        result = []
        for row in rows:
            entry = dict(row)
            entry["detail"] = json.loads(entry.pop("detail"))
            result.append(entry)
        return result


def _count_items(kind: str, payload: Any) -> int:
    if payload is None:
        return 0
    if kind == "anti_dumping":
        return len(payload.get("definitive", [])) + len(payload.get("provisional", []))
    if kind == "surveillance":
        return sum(len(doc.get("items", [])) for doc in payload)
    return len(payload)


def _item_keys(kind: str, payload: Any) -> dict[str, tuple[str, list[str]]]:
    """Değişiklik defteri için satır anahtarı -> (içerik özeti, GTİP kodları)."""
    keys: dict[str, tuple[str, list[str]]] = {}
    if not payload:
        return keys
    if kind == "anti_dumping":
        for row in payload.get("definitive", []):
            codes = extract_codes(row.get("gtip", ""))
            keys[f"{row.get('file_no')}|{row.get('country')}|{' '.join(codes)[:80]}"] = (
                f"{row.get('product')} – {row.get('rate')} ({row.get('communique')}, bitiş {row.get('expires')})",
                codes,
            )
        for row in payload.get("provisional", []):
            keys[f"geçici|{row.get('file_no')}|{row.get('country')}"] = (
                f"{row.get('product')} – {row.get('rate')} ({row.get('communique')})",
                extract_codes(row.get("gtip", "")),
            )
    elif kind == "safeguard":
        for row in payload:
            keys[f"{row.get('file_no')}|{row.get('country')}"] = (
                f"{row.get('product')} – {'; '.join(row.get('amounts', []))} (bitiş {row.get('expires')})",
                extract_codes(row.get("gtip", "")),
            )
    elif kind == "surveillance":
        for doc in payload:
            for item in doc.get("items", []):
                code = normalise_code(item.get("gtip", ""))
                keys[f"{doc.get('mevzuat_no')}|{code}"] = (
                    f"{doc.get('title')}: {item.get('gtip')} {item.get('value')} {item.get('unit') or ''}".strip(),
                    [code] if code else [],
                )
    elif kind == "communiques":
        for row in payload:
            keys[row.get("url", row.get("title", ""))] = (row.get("title", ""), [])
    return keys


def diff_payloads(kind: str, previous: Any, current: Any) -> dict[str, Any]:
    before = _item_keys(kind, previous)
    after = _item_keys(kind, current)
    added = sorted(set(after) - set(before))
    removed = sorted(set(before) - set(after))
    modified = sorted(key for key in set(before) & set(after) if before[key][0] != after[key][0])
    return {
        "kind": kind,
        "label": KIND_LABELS.get(kind, kind),
        "added_count": len(added), "removed_count": len(removed), "modified_count": len(modified),
        "total": len(added) + len(removed) + len(modified),
        "added": [{"key": key, "summary": after[key][0], "codes": after[key][1]} for key in added[:50]],
        "removed": [{"key": key, "summary": before[key][0], "codes": before[key][1]} for key in removed[:50]],
        "modified": [{"key": key, "before": before[key][0], "after": after[key][0], "codes": after[key][1]} for key in modified[:50]],
    }


# --------------------------------------------------------------------------- engine

@dataclass
class SyncOutcome:
    kind: str
    ok: bool
    message: str
    item_count: int = 0
    diff: dict[str, Any] | None = None
    source_url: str = ""


class TradeMeasureEngine:
    def __init__(self, data_dir: str | Path | None = None, *, http: httpx.AsyncClient | None = None) -> None:
        self.store = TradeMeasureStore(data_dir)
        self._http = http
        self._lock = asyncio.Lock()
        self.last_sync: dict[str, Any] = {"started_at": None, "finished_at": None, "outcomes": []}
        self.sync_interval = DEFAULT_SYNC_INTERVAL

    # ---- http
    def _client(self) -> httpx.AsyncClient:
        if self._http is None:
            self._http = httpx.AsyncClient(
                timeout=httpx.Timeout(60.0, connect=20.0),
                headers={"User-Agent": USER_AGENT, "Accept-Language": "tr-TR,tr;q=0.9"},
                verify=official_ssl_context(),
                follow_redirects=True,
            )
        return self._http

    async def aclose(self) -> None:
        if self._http is not None:
            await self._http.aclose()
            self._http = None

    # ---- lookup
    def lookup(self, gtip: str, origin_country: str | None = None, *, today: date | None = None) -> TradeMeasureReport:
        """GTİP (4-12 hane) ve menşe için damping/sübvansiyon, korunma ve gözetim kapsamı."""
        today = today or date.today()
        digits = normalise_code(gtip)
        if len(digits) not in {4, 6, 8, 10, 12}:
            raise ValueError("GTİP 4, 6, 8, 10 veya 12 haneli olmalıdır.")
        report = TradeMeasureReport(gtip=digits, origin_country=origin_country, as_of=today.isoformat())
        report.sources = {kind: self.store.metadata(kind) for kind in ("anti_dumping", "safeguard", "surveillance")}
        if len(digits) < 12:
            report.warnings.append("Kod 12 haneden kısa; ön ek eşleşmeleri kapsam adayıdır, alt satıra göre doğrulayın.")
        report.anti_dumping = self._antidumping_hits(digits, origin_country, today)
        report.safeguard = self._safeguard_hits(digits, origin_country, today)
        report.surveillance = self._surveillance_hits(digits, today)
        if origin_country and any(hit.origin_match is None for hit in report.anti_dumping):
            report.warnings.append("Menşe ülke resmî listedeki ülke adıyla eşleştirilemedi; damping satırlarını elle doğrulayın.")
        for kind in ("anti_dumping", "safeguard", "surveillance"):
            if report.sources[kind]["origin"] == "missing":
                report.warnings.append(f"{KIND_LABELS[kind]} verisi yüklü değil.")
        return report

    def _antidumping_hits(self, gtip: str, origin: str | None, today: date) -> list[MeasureHit]:
        data = self.store.load("anti_dumping") or {}
        hits: list[MeasureHit] = []
        for section, label in (("definitive", "kesin"), ("provisional", "geçici")):
            for row in data.get(section, []):
                matched = [code for code in extract_codes(row.get("gtip", "")) if code_matches(code, gtip)]
                if not matched:
                    continue
                kind = (row.get("kind") or "").upper()
                measure_type = "countervailing" if kind.startswith("SK") or "SÜBVANSİYON" in kind else "anti_dumping"
                hits.append(
                    MeasureHit(
                        measure_type=measure_type,
                        matched_code=max(matched, key=len),
                        country=row.get("country", ""),
                        origin_match=_country_matches(origin, row.get("country", "")),
                        rate_text=row.get("rate", ""),
                        unit_value_usd=None,
                        unit=None,
                        product=row.get("product", ""),
                        legal_act=f"İthalatta Haksız Rekabetin Önlenmesine İlişkin Tebliğ {row.get('communique', '').strip()} ({label} önlem)",
                        gazette=f"RG {row.get('rg_date', '')} / {row.get('rg_no', '')}",
                        expires=row.get("expires") or None,
                        status=_status(row.get("expires"), today) if section == "definitive" else "in_force",
                        notes=" | ".join(part for part in (row.get("kind", ""), row.get("notes", "")) if part),
                        source="Ticaret Bakanlığı – Yürürlükteki Damping/Sübvansiyon Önlemleri listesi",
                    )
                )
        return hits

    def _safeguard_hits(self, gtip: str, origin: str | None, today: date) -> list[MeasureHit]:
        data = self.store.load("safeguard") or []
        hits: list[MeasureHit] = []
        for row in data:
            matched = [code for code in extract_codes(row.get("gtip", "")) if code_matches(code, gtip)]
            if not matched:
                continue
            acts = row.get("acts", [])
            decree = next(
                (act for act in reversed(acts) if "Karar" in act.get("kind", "") and "Geçici" not in act.get("kind", "")),
                acts[-1] if acts else {},
            )
            hits.append(
                MeasureHit(
                    measure_type="safeguard",
                    matched_code=max(matched, key=len),
                    country=row.get("country", ""),
                    origin_match=_country_matches(origin, row.get("country", "")),
                    rate_text="; ".join(row.get("amounts", [])),
                    unit_value_usd=None,
                    unit=None,
                    product=row.get("product", "").replace("\n", " "),
                    legal_act=f"Korunma önlemi kararı {decree.get('number', '')} ({row.get('stage', '')})".strip(),
                    gazette=f"RG {decree.get('rg_date', '')} / {decree.get('rg_no', '')}",
                    expires=row.get("expires") or None,
                    status=_status(row.get("expires"), today),
                    notes=(
                        "Tarife kontenjanı (GYÜ/genel) tahsisi var; kontenjan dâhilinde ithalatta önlem uygulanmaz."
                        if any("Kontenjan" in act.get("kind", "") for act in acts)
                        else ""
                    ),
                    source="Ticaret Bakanlığı – Yürürlükte Bulunan Korunma Önlemleri listesi",
                )
            )
        return hits

    def _surveillance_hits(self, gtip: str, today: date) -> list[MeasureHit]:
        data = self.store.load("surveillance") or []
        hits: list[MeasureHit] = []
        for doc in data:
            for item in doc.get("items", []):
                code = normalise_code(item.get("gtip", ""))
                if not code_matches(code, gtip):
                    continue
                hits.append(
                    MeasureHit(
                        measure_type="surveillance",
                        matched_code=code,
                        country=item.get("country") or "Tüm ülkeler",
                        origin_match=True,
                        rate_text=item.get("value", ""),
                        unit_value_usd=_parse_unit_value(item.get("value", "")),
                        unit=item.get("unit") or doc.get("unit"),
                        product=item.get("description", ""),
                        legal_act=doc.get("title", SURVEILLANCE_TITLE),
                        gazette=f"RG {doc.get('rg_date', '')} / {doc.get('rg_no', '')}",
                        expires=None,
                        status="in_force",
                        notes=item.get("note", ""),
                        source=doc.get("url") or MEVZUAT_PUBLIC.format(no=doc.get("mevzuat_no", "")),
                    )
                )
        return hits

    def communiques(self, year: int | None = None) -> list[dict[str, str]]:
        data = self.store.load("communiques") or []
        if year:
            data = [row for row in data if row.get("year") == str(year)]
        return data

    def status(self) -> dict[str, Any]:
        return {
            "datasets": {kind: {"label": KIND_LABELS[kind], **self.store.metadata(kind)} for kind in KINDS},
            "sync_interval_seconds": self.sync_interval,
            "last_sync": self.last_sync,
            "recent_changes": self.store.changes(limit=10),
        }

    # ---- sync
    async def sync(self, *, kinds: Iterable[str] | None = None) -> list[SyncOutcome]:
        wanted = list(kinds or KINDS)
        async with self._lock:
            self.last_sync = {"started_at": datetime.now(UTC).isoformat(timespec="seconds"), "finished_at": None, "outcomes": []}
            outcomes: list[SyncOutcome] = []
            steps: dict[str, Callable[[], Any]] = {
                "anti_dumping": self._sync_antidumping,
                "safeguard": self._sync_safeguard,
                "surveillance": self._sync_surveillance,
                "communiques": self._sync_communiques,
            }
            for kind in wanted:
                try:
                    outcome = await steps[kind]()
                except Exception as exc:  # noqa: BLE001 – her kaynak bağımsız raporlanır
                    logger.exception("Trade measure sync failed for %s", kind)
                    outcome = SyncOutcome(kind=kind, ok=False, message=f"{type(exc).__name__}: {exc}"[:400])
                outcomes.append(outcome)
                self.last_sync["outcomes"].append(outcome.__dict__)
            self.last_sync["finished_at"] = datetime.now(UTC).isoformat(timespec="seconds")
            return outcomes

    async def _get(self, url: str, **kwargs: Any) -> httpx.Response:
        response = await self._client().get(url, **kwargs)
        response.raise_for_status()
        return response

    async def _sync_antidumping(self) -> SyncOutcome:
        page = await self._get(ANTIDUMPING_PAGE)
        link = discover_workbook_link(page.text, ANTIDUMPING_PAGE, must_contain=("rürlükteki",))
        if not link:
            raise ValueError("Damping sayfasında 'Yürürlükteki Önlemler' çalışma kitabı bağlantısı bulunamadı.")
        payload = parse_antidumping_workbook((await self._get(link)).content)
        diff = self.store.save("anti_dumping", payload, source_url=link, source_label="Ticaret Bakanlığı – Yürürlükteki Önlemler (damping/sübvansiyon)")
        return SyncOutcome("anti_dumping", True, "güncellendi", _count_items("anti_dumping", payload), diff, link)

    async def _sync_safeguard(self) -> SyncOutcome:
        page = await self._get(SAFEGUARD_PAGE)
        link = discover_workbook_link(page.text, SAFEGUARD_PAGE, must_contain=("rürlükteki",))
        if not link:
            raise ValueError("Korunma önlemleri sayfasında yürürlükteki önlemler çalışma kitabı bulunamadı.")
        payload = parse_safeguard_workbook((await self._get(link)).content)
        diff = self.store.save("safeguard", payload, source_url=link, source_label="Ticaret Bakanlığı – Yürürlükte Bulunan Korunma Önlemleri")
        return SyncOutcome("safeguard", True, "güncellendi", len(payload), diff, link)

    async def _sync_communiques(self) -> SyncOutcome:
        year = datetime.now().year
        entries: list[dict[str, str]] = []
        source = ""
        for candidate in (year, year - 1):
            url = COMMUNIQUE_PAGE.format(year=candidate)
            try:
                page = await self._get(url)
            except httpx.HTTPStatusError:
                continue
            parsed = parse_communique_index(page.text, url, candidate)
            if parsed:
                entries.extend(parsed)
                source = source or url
        if not entries:
            raise ValueError("İthalat Tebliğleri dizin sayfasında tebliğ bağlantısı bulunamadı.")
        diff = self.store.save("communiques", entries, source_url=source, source_label="Ticaret Bakanlığı – İthalat Tebliğleri")
        return SyncOutcome("communiques", True, "güncellendi", len(entries), diff, source)

    async def _mevzuat_search(self, title: str) -> list[dict[str, Any]]:
        client = self._client()
        home = await client.get(MEVZUAT_HOME)
        home.raise_for_status()
        token = ""
        for name, value in client.cookies.items():
            if "Antiforgery" in name:
                token = value
        rows: list[dict[str, Any]] = []
        start = 0
        while True:
            payload = {
                "draw": 1,
                "columns": [{"data": None, "name": "", "searchable": True, "orderable": False, "search": {"value": "", "regex": False}}] * 3,
                "order": [], "start": start, "length": 20, "search": {"value": "", "regex": False},
                "parameters": {
                    "MevzuatTur": "Teblig", "YonetmelikMevzuatTur": "OsmanliKanunu", "AranacakIfade": title,
                    "TamCumle": "false", "AranacakYer": "2", "MevzuatNo": "", "KurumId": "0", "AltKurumId": "0",
                    "BaslangicTarihi": "", "BitisTarihi": "", "antiforgerytoken": token,
                },
            }
            response = await client.post(
                MEVZUAT_DATATABLE,
                json=payload,
                headers={"X-Requested-With": "XMLHttpRequest", "Referer": MEVZUAT_HOME},
            )
            response.raise_for_status()
            body = response.json()
            data = body.get("data", [])
            for item in data:
                name = BeautifulSoup(item.get("mevAdi", ""), "html.parser").get_text(" ", strip=True)
                if SURVEILLANCE_TITLE.lower()[:20] not in name.lower() and "gözetim" not in name.lower():
                    continue
                rows.append(
                    {
                        "mevzuat_no": str(item.get("mevzuatNo", "")), "title": name,
                        "rg_date": item.get("resmiGazeteTarihi", ""), "rg_no": item.get("resmiGazeteSayisi", ""),
                        "url": MEVZUAT_PUBLIC.format(no=item.get("mevzuatNo", "")),
                    }
                )
            start += len(data)
            if not data or start >= int(body.get("recordsTotal", 0)):
                break
            await asyncio.sleep(0.5)
        return rows

    async def _sync_surveillance(self) -> SyncOutcome:
        index = await self._mevzuat_search(SURVEILLANCE_TITLE)
        if not index:
            raise ValueError("mevzuat.gov.tr aramasında gözetim tebliği bulunamadı.")
        docs: list[dict[str, Any]] = []
        fetched = 0
        for entry in index:
            cached = self.store.surveillance_doc(entry["mevzuat_no"])
            if cached and cached.get("rg_date") == entry["rg_date"] and cached.get("items") is not None:
                docs.append(cached)
                continue
            response = await self._get(MEVZUAT_IFRAME.format(no=entry["mevzuat_no"]))
            doc = parse_surveillance_page(response.text, entry)
            self.store.save_surveillance_doc(doc)
            docs.append(doc)
            fetched += 1
            await asyncio.sleep(0.4)
        docs.sort(key=lambda doc: (doc.get("rg_date", "")[-4:], doc.get("rg_date", "")[3:5], doc.get("rg_date", "")[:2]), reverse=True)
        diff = self.store.save("surveillance", docs, source_url=MEVZUAT_HOME, source_label="mevzuat.gov.tr – yürürlükteki gözetim tebliğleri")
        return SyncOutcome("surveillance", True, f"güncellendi ({fetched} yeni metin indirildi)", _count_items("surveillance", docs), diff, MEVZUAT_HOME)

    async def periodic_sync_loop(self, *, initial_delay: float = 90.0) -> None:
        """Günde bir (varsayılan) resmî listeleri yeniler; hata olursa bir sonraki turda yeniden dener."""
        await asyncio.sleep(initial_delay)
        while True:
            try:
                outcomes = await self.sync()
                failed = [outcome.kind for outcome in outcomes if not outcome.ok]
                if failed:
                    logger.warning("Trade measure sync finished with failures: %s", ", ".join(failed))
                else:
                    logger.info("Trade measure sync completed")
            except Exception:  # noqa: BLE001
                logger.exception("Trade measure sync loop crashed; retrying next interval")
            await asyncio.sleep(max(3600, self.sync_interval))


def summary_lines(report: TradeMeasureReport) -> list[str]:
    """Kısa, insan okunur özet (danışman ve e-posta çıktıları için)."""
    lines: list[str] = []
    for hit in report.anti_dumping:
        if hit.origin_match is False:
            continue
        label = "Sübvansiyona karşı önlem" if hit.measure_type == "countervailing" else "Dampinge karşı önlem"
        lines.append(f"{label}: {hit.country} menşeli {hit.product} ({hit.matched_code}) – {hit.rate_text} [{hit.legal_act}, {hit.status}]")
    for hit in report.safeguard:
        lines.append(f"Korunma önlemi: {hit.product} ({hit.matched_code}) – {hit.rate_text} [{hit.legal_act}, bitiş {hit.expires}]")
    for hit in report.surveillance:
        lines.append(f"Gözetim: {hit.product} ({hit.matched_code}) – birim kıymet {hit.rate_text} {hit.unit or ''} [{hit.legal_act}]")
    return lines
