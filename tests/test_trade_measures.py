"""Resmî önlem listeleri: ayrıştırma, eşleme, depo ve değişiklik defteri."""

from __future__ import annotations

import io
import json
import struct
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

import trade_measures as tm

SURVEILLANCE_HTML = """
<html><body><p>İTHALATTA GÖZETİM UYGULANMASINA İLİŞKİN TEBLİĞ (TEBLİĞ NO: 2026/37)</p>
<table>
<tr><td>GTİP</td><td>Eşyanın Tanımı</td><td>Birim Gümrük Kıymeti (ABD Doları/Kg*)</td></tr>
<tr><td>8424.90.80.00.12</td><td>Yangın söndürme tüpü tetikleri</td><td>12</td></tr>
<tr><td>8481.10.05.00.00</td><td>Filtre veya yağlayıcılarla kombine halde olanlar</td><td>30</td></tr>
<tr><td>8481.80.59.00.19</td><td>Diğerleri</td><td>2,5</td></tr>
</table>
<p>* Kg: Brüt ağırlık</p></body></html>
"""


def _antidumping_workbook() -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Kesin Önl- Definitive Measures"
    sheet.append(["T.C. TİCARET BAKANLIĞI"])
    header = ["DOSYA NO", "SEKTÖR", "MADDE İSMİ", "MADDE İSMİ", "G.T.İ.P. ", "ÜLKE", "ÜLKE", "ÖNLEM TEBLİĞ NO", "ÖNLEM RG TARİHİ",
              "ÖNLEM RG NO", "BİLGİLENDİRME RAPORU", "ÖNLEM ORANI (CIF%) / MİKTARI", " ÖNLEM TÜRÜ", "NORMAL SÜRE DOLUM TARİHİ"]
    sheet.append(header)
    sheet.append(["CASE NUMBER", "INDUSTRY"])
    sheet.append(["DMS.221.00.2012", "MM", "Paslanmaz Çelik Borular", "Welded", "7306.40.20.90.00 7306.40.80.90.00", "Çin Halk Cumhuriyeti",
                  "China, P.R", "2026/5", "2026-03-15", "33200", None, "%13,82 - %25,27", "DK", "2031-03-15"])
    sheet.append(["NGS.186.01.2011", "MM", "Granitler", "Granites", "6802.23 6802.93", "Çin Halk Cumhuriyeti", "China", "2012/14",
                  "2012-07-10", "28349", None, "174 $/Ton", "DK", "2017-07-10"])
    prov = book.create_sheet("Geçici Önl.-Prov. Measures")
    prov.append(["DOSYA NO", "SEKTÖR", "MADDE İSMİ", "MADDE İSMİ", "G.T.İ.P. ", "ÜLKE", "ÜLKE", "GEÇİCİ ÖNLEM TEBLİĞ NO", "RG TARİHİ", "RG NO", "ORAN", "TÜR", "AÇILIŞ"])
    prov.append(["CASE NUMBER"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _safeguard_workbook() -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([None, "SIRA", "DOSYA /CASE NO", "MADDE İSMİ", "GTİP", "ÜLKE", "TÜRÜ", "KARAR-TEBLİĞ NO", "R.G. NO", "R.G. TRH", "YAYIM", "SIMGE", "UZATMA", "İLK ÖNLEM", "BİTİŞ", "ÖNLEM MİKTARI"])
    sheet.append([None, "1", "242", "POLYESTER ELYAF (İRAN MENŞELİ)", "5503.20.00.00.00", "İran / Iran", "Soruşturma Açılış", "2023/1", "32162", "2023-04-13", None, None, "III. UZATMA (2023-2026)", "2013-09-21", "2026-09-20", "% 17 (21.9.2023-20.9.2024)"])
    sheet.append([None, None, None, None, None, None, "Önlem Karar", "2023/7615", "32310", "2023-09-15", None, None, None, None, None, "% 16 (21.9.2025-20.9.2026)"])
    sheet.append([None, "2", "248", "DİŞ FIRÇALARI", "9603.21.00.00.19", "Tüm Ülkeler/All Countries", "Önlem Karar", "2024/8147", "32449", "03.02.2024", None, None, "II. UZATMA", "2018-02-03", "2027-02-02", "0,09 ABD Doları/Adet"])
    sheet.append([None, None, None, None, None, None, "Kontenjan Dağıtım Genel ve GYÜ/ Tariff Quate", "2024/5", "32460", "14.02.2024", None, None, None, None, None, "GYU Kont. 6.289.963 Adet"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class ParserTests(unittest.TestCase):
    def test_extract_codes_handles_dotted_and_bare_forms(self):
        self.assertEqual(tm.extract_codes("7306.40.20.90.00 7306.40.80.90.00"), ["730640209000", "730640809000"])
        self.assertEqual(tm.extract_codes("55.13 55.14"), ["5513", "5514"])
        self.assertEqual(tm.extract_codes("7004,7005,7006"), ["7004", "7005", "7006"])
        self.assertEqual(tm.extract_codes("5407. ..."), ["5407"])

    def test_antidumping_workbook(self):
        parsed = tm.parse_antidumping_workbook(_antidumping_workbook())
        self.assertEqual(len(parsed["definitive"]), 2)
        self.assertEqual(parsed["definitive"][0]["rate"], "%13,82 - %25,27")
        self.assertEqual(parsed["definitive"][0]["expires"], "2031-03-15")
        self.assertEqual(parsed["provisional"], [])

    def test_safeguard_workbook(self):
        parsed = tm.parse_safeguard_workbook(_safeguard_workbook())
        self.assertEqual([row["file_no"] for row in parsed], ["242", "248"])
        self.assertEqual(len(parsed[0]["acts"]), 2)
        self.assertEqual(parsed[0]["amounts"][-1], "% 16 (21.9.2025-20.9.2026)")
        self.assertIn("Kontenjan", parsed[1]["acts"][-1]["kind"])

    def test_surveillance_page(self):
        doc = tm.parse_surveillance_page(SURVEILLANCE_HTML, {"mevzuat_no": "46246", "rg_date": "09/04/2026", "rg_no": "33219"})
        self.assertEqual(doc["item_count"], 3)
        self.assertEqual(doc["unit"], "ABD Doları/Kg")
        self.assertEqual(doc["items"][2]["value"], "2,5")
        self.assertIn("2026/37", doc["title"])

    def test_workbook_link_discovery_and_communique_index(self):
        html = '<a href="/data/abc/Y&#252;r&#252;rl&#252;kteki &#214;nlemler 13.07.2026.xlsx">Yürürlükteki Önlemler</a><a href="/x.pdf">Diğer</a>'
        link = tm.discover_workbook_link(html, "https://ticaret.gov.tr/ithalat/damping", must_contain=("rürlükteki",))
        self.assertEqual(link, "https://ticaret.gov.tr/data/abc/Yürürlükteki%20Önlemler%2013.07.2026.xlsx")
        index = tm.parse_communique_index(
            '<a href="https://www.resmigazete.gov.tr/eskiler/2025/12/20251231M3-23.pdf">Askıya Alma Sistemine İlişkin Tebliğ (İthalat: 2026/18)</a>'
            '<a href="/ithalat/mevzuat">İthalat Mevzuatı</a>',
            "https://ticaret.gov.tr/ithalat", 2026,
        )
        self.assertEqual(index[0]["number"], "2026/18")
        self.assertEqual(index[0]["gazette"], "20251231")
        self.assertEqual(len(index), 1)

    def test_four_digit_positions_and_text_only_tables(self):
        table = '<table><tr><td>G.T.P.</td><td>Eşyanın Tanımı</td><td>Birim CIF Kıymet (ABD Doları/m2)</td></tr><tr><td>68.09</td><td>Alçı levhalar</td><td>5</td></tr></table>'
        doc = tm.parse_surveillance_page(table, {})
        self.assertEqual(doc["items"][0]["gtip"], "68.09")
        self.assertEqual(doc["unit"], "ABD Doları/m2")
        text = (
            "<p>MADDE 1- (1) Bu Tebliğ ... G.T.İ.P. Eşyanın Tanımı Birim Gümrük Kıymeti (ABD Doları/Kg*) "
            "8481.10.05.00.00 Filtre veya yağlayıcılarla kombine halde olanlar 30 8481.80.99.00.11 Yangın hidrantları 5 "
            "* Kg: brüt ağırlık Gözetim uygulaması MADDE 2- (1) ...</p>"
        )
        doc = tm.parse_surveillance_page(text, {})
        self.assertEqual([item["gtip"] for item in doc["items"]], ["8481.10.05.00.00", "8481.80.99.00.11"])
        self.assertEqual(doc["items"][1]["value"], "5")
        self.assertEqual(doc["parser_version"], tm.PARSER_VERSION)

    def test_spaced_gtip_header_and_cif_value_column(self):
        # Bazı tebliğlerde başlık "G.T.İ .P." biçiminde boşluklu ve sütun adı "CIF Kıymet".
        text = (
            "<p>(1) G.T.İ .P. Eşyanın Tanımı CIF Kıymet (ABD Doları/Ton*) "
            "4820.30.00.00.00 Klasörler, ciltler ve dosya gömlekleri 3.600 "
            "* Ton: Brüt ağırlık Gözetim uygulaması MADDE 2 – (1) ...</p>"
        )
        doc = tm.parse_surveillance_page(text, {})
        self.assertEqual([item["gtip"] for item in doc["items"]], ["4820.30.00.00.00"])
        self.assertEqual(doc["items"][0]["value"], "3.600")

    def test_footnote_marker_after_the_code_is_not_part_of_it(self):
        # Resmî tabloda kod dipnot işaretiyle bitebilir: "4820.30.00.00.00+".
        text = (
            "<p>G.T.İ.P. Eşyanın Tanımı Birim Gümrük Kıymeti (ABD Doları/Kg) "
            "4820.30.00.00.00+ Klasörler ve dosya gömlekleri 3.600 "
            "7306.40.20.90.00* Paslanmaz çelik borular 2.100 "
            "* Kg: brüt ağırlık Gözetim uygulaması MADDE 2</p>"
        )
        doc = tm.parse_surveillance_page(text, {})
        self.assertEqual(
            [item["gtip"] for item in doc["items"]],
            ["4820.30.00.00.00", "7306.40.20.90.00"],
        )
        self.assertEqual([item["value"] for item in doc["items"]], ["3.600", "2.100"])

    def test_scope_written_only_in_the_madde_2_sentence(self):
        # Tablosu olmayan eski tebliğlerde kapsam doğrudan MADDE 2 cümlesindedir.
        text = (
            "<p>Gözetim uygulaması Madde 2 — Gözetim uygulaması 6802.23, 6802.93 ve 6802.99 gümrük tarife "
            "pozisyonlarında yer alan eşyanın CIF kıymeti 500 ABD Doları/ton (brüt ağırlık)'un altında "
            "olanlarının ithalatında ülke ayrımı gözetilmeksizin yapılacaktır.</p>"
        )
        doc = tm.parse_surveillance_page(text, {})
        self.assertEqual([item["gtip"] for item in doc["items"]], ["6802.23", "6802.93", "6802.99"])
        self.assertEqual(doc["items"][0]["value"], "500")
        self.assertEqual(doc["unit"], "ABD Doları/ton")

    def test_single_code_prose_scope_keeps_the_quoted_product_name(self):
        text = (
            '<p>Gözetim uygulaması 3802.90.00.90.13 gümrük tarife istatistik pozisyonlu '
            '“Ağartma toprağı-Asit aktivasyonlu killer”in ithalatında ülke ayrımı '
            'gözetilmeksizin yapılacaktır.</p>'
        )
        doc = tm.parse_surveillance_page(text, {})
        self.assertEqual([item["gtip"] for item in doc["items"]], ["3802.90.00.90.13"])
        self.assertEqual(doc["items"][0]["description"], "Ağartma toprağı-Asit aktivasyonlu killer")

    def test_prose_fallback_does_not_fire_when_a_table_exists(self):
        table = (
            '<table><tr><td>G.T.İ.P.</td><td>Eşyanın Tanımı</td><td>Birim Gümrük Kıymeti (ABD Doları/Kg)</td></tr>'
            '<tr><td>7306.40.20.90.00</td><td>Borular</td><td>2</td></tr></table>'
            "<p>Gözetim uygulaması 6802.23 gümrük tarife pozisyonlarında yer alan eşyanın</p>"
        )
        doc = tm.parse_surveillance_page(table, {})
        self.assertEqual([item["gtip"] for item in doc["items"]], ["7306.40.20.90.00"])

    def test_quota_docx_and_discovery(self):
        import zipfile
        ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        def cell(text):
            return f'<w:tc><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:tc>'
        rows = [
            ["G.T.İ.P.", "Madde İsmi", "Kontenjan Miktarı (Ton)", "Kontenjan Dönemi", "Uygulanacak Gümrük Vergisi (%)"],
            ["0702.00.00.00.00", "Domates", "10.000", "1/1-31/12", "0"],
            ["0805.10", "Portakal", "5.000", "1/1-31/12", "0"],
        ]
        body = "".join("<w:tr>" + "".join(cell(c) for c in r) + "</w:tr>" for r in rows)
        xml = f'<?xml version="1.0"?><w:document xmlns:w="{ns}"><w:body><w:tbl>{body}</w:tbl></w:body></w:document>'
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("word/document.xml", xml)
        doc = tm.parse_quota_document(buffer.getvalue(), "Fas Karar.docx", {"title": "Fas Krallığı Menşeli Bazı Tarım Ürünleri İthalatında Tarife Kontenjanı Uygulanması Hakkında Karar", "url": "https://ticaret.gov.tr/x/Fas%20Karar.docx"})
        self.assertEqual(doc["origin"], "Fas Krallığı")
        self.assertEqual(doc["item_count"], 2)
        self.assertEqual(doc["items"][0]["quantity"], "10.000")
        self.assertEqual(doc["items"][0]["duty_rate"], "0")
        self.assertEqual(doc["items"][1]["period"], "1/1-31/12")
        links = tm.discover_quota_documents('<a href="/data/abc/Fas Karar.doc">Fas Krallığı Menşeli Bazı Tarım Ürünleri İthalatında Tarife Kontenjanı</a><a href="/data/abc/x.pdf">Kontenjan pdf</a>', "https://ticaret.gov.tr/ithalat")
        self.assertEqual(len(links), 1)
        self.assertEqual(links[0]["url"], "https://ticaret.gov.tr/data/abc/Fas%20Karar.doc")

    def test_legacy_doc_piece_table_and_rows(self):
        """Word 97-2003 eki: parça tablosu çözülür, hücreler satırlara ayrılır (antiword gerekmez)."""
        text = "Tarife Kontenjanı Kod No\x07G.T.İ.P.\x07Madde İsmi\x07Tarife Kontenjanı Dönemi\x07\x07ARN01\x0704.06\x07Peynir\x0701.01-31.12\x07\x07"
        encoded = text.encode("utf-16-le")
        document = bytearray(0x400)
        start = 0x300
        document[start : start + len(encoded)] = encoded
        piece_table = struct.pack("<II", 0, len(text)) + struct.pack("<HIH", 0, start, 0)
        clx = b"\x02" + struct.pack("<I", len(piece_table)) + piece_table
        struct.pack_into("<II", document, 0x01A2, 0, len(clx))
        decoded = tm._decode_doc_pieces(bytes(document), clx)
        self.assertEqual(
            tm._doc_rows_from_text(decoded),
            [["Tarife Kontenjanı Kod No", "G.T.İ.P.", "Madde İsmi", "Tarife Kontenjanı Dönemi"], ["ARN01", "04.06", "Peynir", "01.01-31.12"]],
        )
        items = tm._quota_rows_from_cells(tm._doc_rows_from_text(decoded))
        self.assertEqual(items[0]["gtip"], "04.06")
        self.assertEqual(items[0]["quota_code"], "ARN01")
        self.assertEqual(items[0]["period"], "01.01-31.12")
        self.assertEqual(tm._doc_table_rows(b"not an OLE document"), [])
        self.assertEqual(tm._decode_doc_pieces(b"", b""), "")

    def test_quota_header_ignores_merged_title_and_realigns_short_rows(self):
        """Başlık hücresine karışan künye ("Tarihi") dönem sütununu kapmamalı, kısa satırlar hizalanmalı."""
        merged_title = "BİRLEŞİK ARAP EMİRLİKLERİ MENŞELİ ... Resmî Gazete'nin Tarihi Sayısı Tarife Kontenjanı Kod No"
        rows = [
            [merged_title, "GTİP", "Eşyanın Tanımı", "Tarife Kontenjanı Dönemi", "Azami Miktar", "Tahsisat Yöntemi"],
            ["BAET001", "0407.21", "Tavuk yumurtaları", "01.01-31.12"],
        ]
        item = tm._quota_rows_from_cells(rows)[0]
        self.assertEqual(item["quota_code"], "BAET001")
        self.assertEqual(item["period"], "01.01-31.12")
        self.assertNotIn("quantity", item)
        self.assertEqual(item["description"], "Tavuk yumurtaları")

        shifted = [
            ["G.T.İ.P.", "Madde İsmi", "Kontenjan Miktarı", "Vergi"],
            ["1", "0406.90.99.00.11", "100 ton", "0"],   # başlıkta olmayan sıra numarası
            ["01.02", "8.000 ton", "0"],                  # birleştirilmiş açıklama hücresi
            ["0102.29", "Diğerleri", "2.260 ton", "0"],
        ]
        parsed = tm._quota_rows_from_cells(shifted)
        self.assertEqual([row["gtip"] for row in parsed], ["0406.90.99.00.11", "01.02", "0102.29"])
        self.assertEqual([row["quantity"] for row in parsed], ["100 ton", "8.000 ton", "2.260 ton"])
        self.assertEqual([row["duty_rate"] for row in parsed], ["0", "0", "0"])
        self.assertEqual(parsed[2]["description"], "Diğerleri")


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        seed = Path(self.tmp) / "seed"
        seed.mkdir()
        (seed / "antidumping_measures.json").write_text(json.dumps(tm.parse_antidumping_workbook(_antidumping_workbook())), encoding="utf-8")
        (seed / "safeguard_measures.json").write_text(json.dumps(tm.parse_safeguard_workbook(_safeguard_workbook())), encoding="utf-8")
        (seed / "surveillance_measures.json").write_text(
            json.dumps([tm.parse_surveillance_page(SURVEILLANCE_HTML, {"mevzuat_no": "46246", "rg_date": "09/04/2026", "rg_no": "33219", "title": "Gözetim 2026/37"})]),
            encoding="utf-8",
        )
        (seed / "agricultural_quotas.json").write_text(json.dumps([{"title": "Fas Krallığı Menşeli Bazı Tarım Ürünleri İthalatında Tarife Kontenjanı Uygulanması Hakkında Karar", "origin": "Fas Krallığı", "url": "https://ticaret.gov.tr/x/Fas.docx", "items": [{"gtip": "0805.10", "description": "Portakal", "quantity": "5.000 Ton", "period": "1/1-31/12", "duty_rate": "0"}]}]), encoding="utf-8")
        self.engine = tm.TradeMeasureEngine(self.tmp)
        self.engine.store = tm.TradeMeasureStore(self.tmp, seed_dir=seed)

    def test_lookup_matches_prefix_and_origin(self):
        report = self.engine.lookup("730640209000", "Çin", today=date(2026, 9, 6))
        self.assertEqual(len(report.anti_dumping), 1)
        hit = report.anti_dumping[0]
        self.assertTrue(hit.origin_match)
        self.assertEqual(hit.status, "in_force")
        self.assertEqual(hit.matched_code, "730640209000")
        other = self.engine.lookup("730640209000", "Almanya", today=date(2026, 9, 6))
        self.assertFalse(other.anti_dumping[0].origin_match)
        self.assertEqual(other.applicable_anti_dumping, [])

    def test_taiwan_measures_do_not_leak_into_china_lookups(self):
        store = self.engine.store
        rows = store.load("anti_dumping")
        rows["definitive"].append({
            "file_no": "NGS.152.04.2024", "sector": "TK", "product": "Polyester Elyaf", "product_en": "",
            "gtip": "5503.20.00.00.00", "country": "Çin Tayvanı", "country_en": "Chinese Taipei",
            "communique": "2026/1", "rg_date": "2026-01-30", "rg_no": "33153", "rate": "%3,2-%12",
            "kind": "DK", "expires": "2031-01-30", "notes": "",
        })
        store.save("anti_dumping", rows, source_url="https://example.test/a.xlsx", source_label="test")
        taiwan = [hit for hit in self.engine.lookup("550320000000", "Tayvan").anti_dumping if hit.country == "Çin Tayvanı"]
        china = [hit for hit in self.engine.lookup("550320000000", "Çin").anti_dumping if hit.country == "Çin Tayvanı"]
        self.assertTrue(taiwan and taiwan[0].origin_match)
        self.assertTrue(china and china[0].origin_match is False)

    def test_expired_measure_flagged(self):
        report = self.engine.lookup("680223000000", "Çin", today=date(2026, 9, 6))
        self.assertEqual(report.anti_dumping[0].status, "expired")

    def test_safeguard_and_surveillance(self):
        report = self.engine.lookup("960321000019", "Vietnam", today=date(2026, 9, 6))
        self.assertEqual(len(report.safeguard), 1)
        self.assertTrue(report.safeguard[0].origin_match)
        self.assertIn("kontenjan", report.safeguard[0].notes.lower())
        surveillance = self.engine.lookup("848180590019", None)
        self.assertEqual(surveillance.surveillance_unit_value, 2.5)
        self.assertEqual(surveillance.surveillance[0].unit, "ABD Doları/Kg")
        lines = tm.summary_lines(surveillance)
        self.assertTrue(lines and lines[0].startswith("Gözetim"))

    def test_quota_hits_match_origin(self):
        report = self.engine.lookup("080510220000", "Fas")
        self.assertEqual(len(report.tariff_quota), 1)
        self.assertTrue(report.tariff_quota[0].origin_match)
        self.assertIn("5.000 Ton", report.tariff_quota[0].rate_text)
        other = self.engine.lookup("080510220000", "Mısır")
        self.assertFalse(other.tariff_quota[0].origin_match)
        self.assertTrue(any(line.startswith("Tarife kontenjanı") for line in tm.summary_lines(report)))

    def test_short_code_matches_longer_official_rows(self):
        report = self.engine.lookup("848180", None)
        self.assertEqual(len(report.surveillance), 1)
        self.assertTrue(report.warnings)

    def test_store_save_records_changes(self):
        store = self.engine.store
        previous = store.load("safeguard")
        updated = json.loads(json.dumps(previous))
        updated[0]["expires"] = "2029-09-20"
        updated.append({"seq": 3, "file_no": "260", "product": "ETİL ASETAT", "gtip": "2915.31.00.00.00", "country": "Tüm Ülkeler", "stage": "İLK", "original_start": "", "expires": "2028-06-22", "amounts": ["% 10"], "acts": []})
        diff = store.save("safeguard", updated, source_url="https://example.test/k.xlsx", source_label="test")
        self.assertEqual((diff["added_count"], diff["removed_count"], diff["modified_count"]), (1, 0, 1))
        self.assertEqual(store.metadata("safeguard")["origin"], "synced")
        changes = store.changes("safeguard")
        self.assertEqual(len(changes), 1)
        self.assertEqual(changes[0]["detail"]["added"][0]["codes"], ["291531000000"])
        # Yeni depo örneği aynı veriyi sqlite'tan okur.
        reopened = tm.TradeMeasureStore(self.tmp, seed_dir=Path(self.tmp) / "seed")
        self.assertEqual(len(reopened.load("safeguard")), 3)

    def test_status_reports_seed_origin(self):
        status = self.engine.status()
        self.assertEqual(status["datasets"]["anti_dumping"]["origin"], "seed")
        self.assertEqual(status["datasets"]["communiques"]["origin"], "missing")
        self.assertEqual(status["datasets"]["tariff_quota"]["origin"], "seed")


class OfficialSeedTests(unittest.TestCase):
    """Depodaki resmî tohum dosyaları ve gerçek fihrist sayfası ayrıştırması."""

    def test_real_iframe_page_parses(self):
        html_text = (Path(__file__).parent / "fixtures" / "surveillance_iframe_sample.htm").read_text(encoding="utf-8", errors="ignore")
        doc = tm.parse_surveillance_page(html_text, {"mevzuat_no": "46246", "rg_date": "09/04/2026", "rg_no": "33219"})
        self.assertGreaterEqual(doc["item_count"], 1)
        self.assertTrue(all(tm.normalise_code(item["gtip"]) for item in doc["items"]))

    def test_seed_files_load_with_default_store(self):
        engine = tm.TradeMeasureEngine(tempfile.mkdtemp())
        status = engine.status()["datasets"]
        self.assertGreater(status["anti_dumping"]["item_count"], 200)
        self.assertGreater(status["safeguard"]["item_count"], 5)
        self.assertGreater(status["surveillance"]["item_count"], 500)
        self.assertGreater(status["communiques"]["item_count"], 10)
        report = engine.lookup("730640209000", "Çin", today=date(2026, 9, 6))
        self.assertTrue(report.applicable_anti_dumping)


if __name__ == "__main__":
    unittest.main()
